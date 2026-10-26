"""
Passo 0 - Preparazione dei dati.

Legge il file Excel ufficiale ISTAT con l'elenco dei comuni italiani
e salva un nome per riga in data/comuni.txt, pronto per l'addestramento.
"""

import random
from pathlib import Path

from openpyxl import load_workbook

CARTELLA_DATI = Path("data")
FILE_ISTAT = CARTELLA_DATI / "Elenco-comuni-italiani.xlsx"
FILE_OUTPUT = CARTELLA_DATI / "comuni.txt"
COLONNA_NOME = "denominazione in italiano"


def trova_colonna(riga: tuple) -> int | None:
    """
    Cerca in una riga la cella di intestazione del nome del comune.
    Restituisce la posizione della colonna, oppure None se non c'è.
    """
    for indice, cella in enumerate(riga):
        if isinstance(cella, str):
            # " ".join(split()) toglie a capo e spazi doppi nell'intestazione
            testo = " ".join(cella.lower().split())
            if COLONNA_NOME in testo:
                return indice
    return None


def carica_comuni(percorso: Path) -> list[str]:
    cartella_excel = load_workbook(percorso, read_only=True)
    foglio = cartella_excel.active

    indice_colonna = None
    nomi = set()  # il set elimina in automatico i duplicati

    for riga in foglio.iter_rows(values_only=True):
        # Finché non troviamo l'intestazione, cerchiamo solo quella
        if indice_colonna is None:
            indice_colonna = trova_colonna(riga)
            continue

        if indice_colonna >= len(riga):
            continue

        valore = riga[indice_colonna]
        if isinstance(valore, str) and valore.strip():
            nomi.add(valore.strip().lower())

    cartella_excel.close()

    if indice_colonna is None:
        raise KeyError(f"Colonna '{COLONNA_NOME}' non trovata nel file")

    return sorted(nomi)


def stampa_statistiche(nomi: list[str]) -> None:
    lunghezze = [len(n) for n in nomi]
    vocabolario = sorted(set("".join(nomi)))

    print(f"Comuni caricati:   {len(nomi)}")
    print(f"Lunghezza minima:  {min(lunghezze)}")
    print(f"Lunghezza massima: {max(lunghezze)}")
    print(f"Lunghezza media:   {sum(lunghezze) / len(lunghezze):.1f}")
    print(f"Caratteri diversi: {len(vocabolario)}")
    print(f"Vocabolario:       {vocabolario}")
    print(f"Esempi casuali:    {random.sample(nomi, 8)}")


def main() -> None:
    nomi = carica_comuni(FILE_ISTAT)
    FILE_OUTPUT.write_text("\n".join(nomi), encoding="utf-8")
    stampa_statistiche(nomi)
    print(f"\nSalvato in {FILE_OUTPUT}")


if __name__ == "__main__":
    main()